import re, subprocess, copy, json, sys, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D = "/home/claude/nivotime/final/"
PURPLE = RGBColor(0x3B, 0x3A, 0x73); GREEN = RGBColor(0x5E, 0x8A, 0x2E); GRAY = RGBColor(0x4A, 0x4D, 0x5E)
FONT = "Times New Roman"
TEXT_W = 6.27  # inches, A4 with 1in margins

# ---------- 1. reference doc ----------
subprocess.run(f"pandoc -o {D}ref.docx --print-default-data-file reference.docx", shell=True, check=True)
ref = Document(D + "ref.docx")
def S(name):
    for st in ref.styles:
        if st.name == name: return st
    raise KeyError(name)
def setfont(st, size=None, bold=None, italic=None, color=None, name=FONT):
    f = st.font; f.name = name
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ["w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"]: rf.set(qn(a), name)
    for a in ["w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"]:
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color
for name in ["Normal", "Body Text", "First Paragraph", "Compact"]:
    try: st = S(name)
    except KeyError: continue
    setfont(st, 12, color=RGBColor(0, 0, 0))
    pf = st.paragraph_format; pf.line_spacing = 1.5; pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6); pf.space_before = Pt(0)
cst = S("Compact"); cst.paragraph_format.space_after = Pt(3)
for lvl, size, before, after in [(1, 17, 0, 10), (2, 13.5, 14, 6), (3, 12, 10, 4)]:
    st = S(f"Heading {lvl}")
    setfont(st, size, bold=True, italic=False, color=PURPLE if lvl < 3 else GREEN)
    pf = st.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.keep_with_next = True; pf.line_spacing = 1.15; pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
for name in ["Image Caption", "Table Caption", "Caption"]:
    try:
        st = S(name); setfont(st, 10, italic=True, color=GRAY)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; st.paragraph_format.space_after = Pt(10)
        st.paragraph_format.line_spacing = 1.0
    except KeyError: pass
try:
    st = S("Captioned Figure"); st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
except KeyError: pass
bt = S("Block Text"); setfont(bt, 12, italic=False, color=PURPLE)
bt.paragraph_format.left_indent = Inches(0.35); bt.paragraph_format.right_indent = Inches(0.3)
bt.paragraph_format.space_before = Pt(6); bt.paragraph_format.space_after = Pt(8); bt.paragraph_format.line_spacing = 1.3
ppr = bt.element.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
for k, v in {"w:val": "single", "w:sz": "24", "w:space": "10", "w:color": "87B549"}.items(): l.set(qn(k), v)
bdr.append(l); ppr.append(bdr)
shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F7EC"); ppr.append(shd)
ref.save(D + "ref.docx")

# ---------- 2. markdown ----------
parts = ["report_part1.md", "report_part2.md", "report_part3.md", "report_part4.md", "report_refs.md"]
md = "\n\n".join(open(D + p).read() for p in parts)
md = md.replace("[[PB]]", "", 1)
md = md.replace("[[PB]]", '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```')
open(D + "report_full.md", "w").write(md)
subprocess.run(f"cd {D} && LANG=en_US.UTF-8 pandoc report_full.md -f markdown-auto_identifiers+gfm_auto_identifiers --reference-doc=ref.docx -o body.docx",
               shell=True, check=True)

# ---------- 3. post-process ----------
doc = Document(D + "body.docx")
body = doc.element.body

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)

def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "C9CCD6"); b.append(e)
    tblPr.append(b)

for tbl in doc.tables:
    rows = tbl.rows; ncol = len(tbl.columns)
    # auto column widths from content
    lens = [0]*ncol
    for r in rows:
        for j, c in enumerate(r.cells[:ncol]):
            t = c.text.strip()
            longest_word = max([len(w) for w in t.split()] + [1])
            lens[j] = max(lens[j], min(len(t), 70)**0.75 + longest_word*0.35)
    lens = [max(x, 3.5) for x in lens]
    minw = []
    for j in range(ncol):
        lw = max([max([len(w) for w in r.cells[j].text.split()] + [1]) for r in rows])
        minw.append(min(lw*0.085 + 0.22, 1.6))
    tot = sum(lens); widths = [TEXT_W*x/tot for x in lens]
    for _ in range(10):
        short = [j for j in range(ncol) if widths[j] < minw[j]]
        if not short: break
        for j in short: widths[j] = minw[j]
        fixed = sum(widths[j] for j in range(ncol) if widths[j] <= minw[j] + 1e-9)
        free = [j for j in range(ncol) if widths[j] > minw[j] + 1e-9]
        ft = sum(lens[j] for j in free)
        for j in free: widths[j] = max(minw[j], (TEXT_W - fixed)*lens[j]/ft)
    tblPr = tbl._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(int(TEXT_W*1440))); tw.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j]*1440)))
    set_borders(tbl)
    empty_header = all(not c.text.strip() for c in rows[0].cells)
    for i, r in enumerate(rows):
        for j, c in enumerate(r.cells[:ncol]):
            c.width = Inches(widths[j])
            tcPr = c._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side in ["top", "bottom"]:
                m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), "40"); m.set(qn("w:type"), "dxa"); mar.append(m)
            tcPr.append(mar)
            if i == 0 and not empty_header: shade(c, "3B3A73")
            elif i % 2 == 0: shade(c, "F4F5F9")
            for p in c.paragraphs:
                pf = p.paragraph_format; pf.line_spacing = 1.05; pf.space_after = Pt(1); pf.space_before = Pt(1)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(9.5); run.font.name = FONT
                    if i == 0 and not empty_header:
                        run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    if empty_header:  # 2-col key/value tables (CEO summary): drop empty header row, bold first col
        tbl._tbl.remove(rows[0]._tr)
        for r in tbl.rows:
            shade(r.cells[0], "ECEBF7")
            for p in r.cells[0].paragraphs:
                for run in p.runs: run.font.bold = True; run.font.color.rgb = PURPLE
    # repeat header row
    if not empty_header:
        trPr = tbl.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for r in tbl.rows:
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)

for tbl in doc.tables:
    nxt = tbl._tbl.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        from docx.text.paragraph import Paragraph
        Paragraph(nxt, None).paragraph_format.space_before = Pt(8)
# list paragraphs: tighter spacing
for p in doc.paragraphs:
    if p._p.pPr is not None and p._p.pPr.numPr is not None:
        p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.35

# Section-divider headings ("SECTION A — ...") styled as banners
for p in doc.paragraphs:
    if p.style.name == "Heading 1" and p.text.startswith("SECTION "):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pPr = p._p.get_or_add_pPr(); s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), "3B3A73"); pPr.append(s)
        for r in p.runs: r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(14)

for p in doc.paragraphs:
    t = p.text.strip()
    if t and len(t) < 70 and p.runs and all(r.bold for r in p.runs if r.text.strip()) and p.style.name in ("Body Text","First Paragraph","Normal"):
        p.paragraph_format.keep_with_next = True
in_refs = False
for p in doc.paragraphs:
    if p.style.name == "Heading 1": in_refs = (p.text.strip() == "References")
    elif in_refs: p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing = 1.15

# ---------- section: A4, margins, borders, header/footer ----------
sec = doc.sections[0]
sec.page_width = Emu(7560310); sec.page_height = Emu(10692130)
for side in ["left_margin", "right_margin"]: setattr(sec, side, Inches(1))
sec.top_margin = Inches(1.05); sec.bottom_margin = Inches(0.95)
sec.header_distance = Inches(0.45); sec.footer_distance = Inches(0.4)
sectPr = sec._sectPr
pgB = OxmlElement("w:pgBorders"); pgB.set(qn("w:offsetFrom"), "page")
for edge in ["top", "left", "bottom", "right"]:
    e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "thinThickSmallGap"); e.set(qn("w:sz"), "18")
    e.set(qn("w:space"), "24"); e.set(qn("w:color"), "3B3A73"); pgB.append(e)
# pgBorders must come after pgMar in sectPr order
pgMar = sectPr.find(qn("w:pgMar")); pgMar.addnext(pgB)
sec.different_first_page_header_footer = True

def add_field(run, instr):
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in [f1, it, f2, t, f3]: run._r.append(el)

hp = sec.header.paragraphs[0]; hp.text = ""
hp.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
hp.add_run().add_picture(D + "logo_new.png", height=Inches(0.3))
r = hp.add_run("\tBenefits Intelligence Strategy  |  Live Company Project")
r.font.size = Pt(9); r.font.color.rgb = GRAY; r.font.name = FONT; r.font.italic = True
pPr = hp._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
for k, v in {"w:val": "single", "w:sz": "8", "w:space": "4", "w:color": "87B549"}.items(): b.set(qn(k), v)
bdr.append(b); pPr.append(bdr)
fp = sec.footer.paragraphs[0]; fp.text = ""
fp.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
r = fp.add_run("NivoTime Pvt. Ltd.  ·  WeSchool PGDM (E-Business) 2025–27"); r.font.size = Pt(8.5); r.font.color.rgb = GRAY; r.font.name = FONT
r = fp.add_run("\tPage "); r.font.size = Pt(9); r.font.color.rgb = PURPLE; r.font.name = FONT; r.font.bold = True
def frun(el):
    r = fp.add_run(); r.font.size = Pt(9); r.font.color.rgb = PURPLE; r.font.bold = True; r.font.name = FONT; r._r.append(el); return r
e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "begin"); frun(e)
e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = " PAGE "; frun(e)
e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "separate"); frun(e)
e = OxmlElement("w:t"); e.text = "1"; frun(e)
e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); frun(e)

# ---------- cover + TOC ----------
first = body[0]
def P(text="", size=12, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, spacing=1.0, caps_spacing=None):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = spacing
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.name = FONT
        if color is not None: r.font.color.rgb = color
        if caps_spacing:
            rpr = r._r.get_or_add_rPr(); sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), str(caps_spacing)); rpr.append(sp)
    first.addprevious(p._p); return p
def rule(color="87B549", sz="18"):
    p = P(before=4, after=10); pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": sz, "w:space": "1", "w:color": color}.items(): b.set(qn(k), v)
    bdr.append(b); pPr.append(bdr)
def pagebreak():
    p = doc.add_paragraph(); p.add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE); first.addprevious(p._p)

P(before=36)
lp = P(); lp.add_run().add_picture(D + "logo_new.png", width=Inches(2.9))
P(before=30)
P("LIVE COMPANY PROJECT REPORT", 11, bold=True, color=GREEN, caps_spacing=60, after=10)
P("From Services Breadth to", 26, bold=True, color=PURPLE, spacing=1.05)
P("Benefits Intelligence", 26, bold=True, color=PURPLE, spacing=1.05, after=10)
P("An AI-led product, customer-targeting and go-to-market strategy", 14, italic=True, color=GRAY)
P("for NivoTime Private Limited", 14, italic=True, color=GRAY, after=4)
rule()
P("Strategy  ·  Product & AI  ·  Customer Targeting  ·  Go-to-Market  ·  Sales Engine  ·  Financials", 10.5, color=PURPLE, after=60)
P("Prepared by", 10, color=GRAY, after=2)
P("Anvesha", 15, bold=True, color=PURPLE, after=2)
P("PGDM (E-Business), Batch 2025–27", 11.5, after=2)
P("Prin. L.N. Welingkar Institute of Management Development & Research (WeSchool), Mumbai", 11.5, after=18)
P("Industry partner", 10, color=GRAY, after=2)
P("NivoTime Private Limited, Mumbai", 11.5, bold=True, after=18)
P("September 2026", 11.5, color=GRAY)
pagebreak()
tp = P("Table of Contents", 18, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.LEFT, after=8)
rule("3B3A73", "8")
toc_marker = P("[[TOC]]", align=WD_ALIGN_PARAGRAPH.LEFT)
first_h = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
first_h.paragraph_format.page_break_before = True
doc.save(D + "stage.docx")
print("stage saved")
